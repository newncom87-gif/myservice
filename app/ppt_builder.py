"""theme(위치/서식)과 구절 조합 결과를 받아 실제 PPTX 파일을 생성한다.

프론트엔드 프리뷰와 좌표계를 맞추기 위해 모든 위치/크기는 슬라이드 폭·높이에
대한 비율(0~1, xPct/yPct/wPct/hPct)로 표현한다.
"""
from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import backgrounds, bible_data as bd, compose

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
EMU_PER_IN = 914400

ALIGN_MAP = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
OVERLAY_SHAPE_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded-rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
}


def _color(hex_str):
    hex_str = (hex_str or "#FFFFFF").lstrip("#")
    if len(hex_str) != 6:
        hex_str = "FFFFFF"
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


def _add_background(slide, prs, image_path):
    pic = slide.shapes.add_picture(str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)


def _add_overlay(slide, prs, style):
    """배경과 텍스트 사이에 반투명 색 도형을 넣어 가독성을 높인다."""
    if not style or not style.get("enabled"):
        return None
    x = Emu(int(float(style["xPct"]) * prs.slide_width))
    y = Emu(int(float(style["yPct"]) * prs.slide_height))
    w = Emu(int(float(style["wPct"]) * prs.slide_width))
    h = Emu(int(float(style["hPct"]) * prs.slide_height))
    mso_shape = OVERLAY_SHAPE_MAP.get(style.get("shape", "rect"), MSO_SHAPE.RECTANGLE)

    shape = slide.shapes.add_shape(mso_shape, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(style.get("color", "#000000"))
    shape.line.fill.background()
    shape.shadow.inherit = False

    # 테마의 스타일 참조(효과/그림자 등)를 제거해 순수 반투명 색상만 남긴다
    style_el = shape._element.find(qn("p:style"))
    if style_el is not None:
        shape._element.remove(style_el)

    opacity = float(style.get("opacity", 0.35))
    alpha_val = str(int(max(0.0, min(1.0, opacity)) * 100000))
    srgb = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
    etree.SubElement(srgb, qn("a:alpha")).set("val", alpha_val)
    return shape


def _add_textbox(slide, prs, style, text):
    if not style or not text:
        return None
    x = Emu(int(float(style["xPct"]) * prs.slide_width))
    y = Emu(int(float(style["yPct"]) * prs.slide_height))
    w = Emu(int(float(style["wPct"]) * prs.slide_width))
    h = Emu(int(float(style["hPct"]) * prs.slide_height))
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN_MAP.get(style.get("align", "center"), PP_ALIGN.CENTER)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(float(style.get("fontSize", 28)))
        run.font.bold = bool(style.get("bold", False))
        run.font.italic = bool(style.get("italic", False))
        run.font.name = style.get("fontFamily") or "맑은 고딕"
        run.font.color.rgb = _color(style.get("color", "#FFFFFF"))
    return box


def _verse_number_label(v):
    if v["vEnd"] and v["vEnd"] != v["v"]:
        return f"{v['v']}~{v['vEnd']}. "
    return f"{v['v']}. "


# 절 번호 라벨("30. ", "18~19. " 등)의 실제 렌더링 폭을 문자별 상대폭(em, 글자크기
# 대비 비율)으로 근사해 행잉 인덴트 폭을 계산한다. 값은 "맑은 고딕"/"Apple SD Gothic
# Neo"를 브라우저에서 실측(getBoundingClientRect)해 캘리브레이션한 것.
_CHAR_EM_WIDTH = {".": 0.30, " ": 0.313, "~": 0.72}
_DIGIT_EM_WIDTH = 0.58


def _label_indent_pt(label, font_size_pt):
    total_em = sum(_DIGIT_EM_WIDTH if ch.isdigit() else _CHAR_EM_WIDTH.get(ch, 0.5) for ch in label)
    return total_em * font_size_pt


def _verse_key(v):
    """어느 슬라이드에 배치되든 일관된 절 식별자. 프론트엔드 verseKey()와 형식을 맞춰야 한다."""
    return f"{v['book']}:{v['chapter']}:{v['v']}:{v['vEnd']}"


def _add_verse_textbox(slide, prs, style, verses, manual_breaks=None):
    """절 번호(강조색)+본문을 절마다 한 문단으로 넣는다.

    절 사이 간격은 빈 문단이 아니라 space_after 로 조절하고(값 조정 가능),
    왼쪽 정렬일 때는 행잉 인덴트를 적용해 줄바꿈된 본문 둘째 줄이 번호가
    아니라 첫 줄 본문 시작 위치에 맞춰지도록 한다. manual_breaks에 해당 절의
    수동 줄바꿈 텍스트(\n 포함)가 있으면 원문 대신 그것을 쓰고, \n 위치마다
    <a:br/>로 강제 줄바꿈한다(새 문단이 아니라 같은 문단 안에서).
    """
    manual_breaks = manual_breaks or {}
    if not style or not verses:
        return None
    x = Emu(int(float(style["xPct"]) * prs.slide_width))
    y = Emu(int(float(style["yPct"]) * prs.slide_height))
    w = Emu(int(float(style["wPct"]) * prs.slide_width))
    h = Emu(int(float(style["hPct"]) * prs.slide_height))
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    align_key = style.get("align", "left")
    align = ALIGN_MAP.get(align_key, PP_ALIGN.LEFT)
    font_name = style.get("fontFamily") or "맑은 고딕"
    font_size_pt = float(style.get("fontSize", 28))
    font_size = Pt(font_size_pt)
    number_color = _color(style.get("numberColor") or style.get("color", "#FFFFFF"))
    text_color = _color(style.get("color", "#FFFFFF"))
    number_bold = bool(style.get("numberBold", False))
    text_bold = bool(style.get("bold", False))
    verse_spacing_pt = float(style.get("verseSpacing", 10))

    for i, v in enumerate(verses):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i < len(verses) - 1:
            p.space_after = Pt(verse_spacing_pt)
        label = _verse_number_label(v)
        if align_key == "left":
            hang_indent_emu = int(Pt(_label_indent_pt(label, font_size_pt)))
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(hang_indent_emu))
            pPr.set("indent", str(-hang_indent_emu))

        r1 = p.add_run()
        r1.text = label
        r1.font.size = font_size
        r1.font.bold = number_bold
        r1.font.name = font_name
        r1.font.color.rgb = number_color

        text = manual_breaks.get(_verse_key(v), v["text"])
        for j, segment in enumerate(text.split("\n")):
            if j > 0:
                p._p.add_br()
            r2 = p.add_run()
            r2.text = segment
            r2.font.size = font_size
            r2.font.bold = text_bold
            r2.font.name = font_name
            r2.font.color.rgb = text_color
    return box


def build_presentation(spec):
    blocks = spec.get("blocks") or []
    if not blocks:
        raise ValueError("구절 블록이 하나 이상 필요합니다")

    default_group_size = spec.get("defaultGroupSize", 2)
    overrides = spec.get("overrides") or []
    theme = spec.get("theme") or {}
    cover_theme = theme.get("cover") or {}
    body_theme = theme.get("body") or {}
    global_bg = spec.get("background")
    slide_backgrounds = spec.get("slideBackgrounds") or {}
    manual_breaks = spec.get("manualBreaks") or {}

    flat_verses, block_refs = compose.resolve_blocks(blocks)
    slides_data = compose.compute_slides(flat_verses, default_group_size, overrides)

    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W_IN * EMU_PER_IN))
    prs.slide_height = Emu(int(SLIDE_H_IN * EMU_PER_IN))
    blank_layout = prs.slide_layouts[6]

    def bg_for(key):
        ref = slide_backgrounds.get(str(key)) or global_bg
        return backgrounds.resolve_path(ref) if ref else None

    # 표지 슬라이드
    cover_slide = prs.slides.add_slide(blank_layout)
    cover_bg = bg_for(0)
    if cover_bg:
        _add_background(cover_slide, prs, cover_bg)
    _add_overlay(cover_slide, prs, cover_theme.get("overlay"))
    _add_textbox(cover_slide, prs, cover_theme.get("title"), (cover_theme.get("title") or {}).get("text", ""))
    _add_textbox(cover_slide, prs, cover_theme.get("verseRef"), ", ".join(block_refs))
    presenter_style = cover_theme.get("presenter") or {}
    if presenter_style.get("show", True):
        _add_textbox(cover_slide, prs, presenter_style, presenter_style.get("text", ""))

    # 본문(구절) 슬라이드
    for i, verses in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank_layout)
        bg = bg_for(i)
        if bg:
            _add_background(slide, prs, bg)
        _add_overlay(slide, prs, body_theme.get("overlay"))
        _add_verse_textbox(slide, prs, body_theme.get("verseText"), verses, manual_breaks)
        ref_title_style = body_theme.get("refTitle") or {}
        if ref_title_style.get("show", True):
            _add_textbox(slide, prs, ref_title_style, bd.format_slide_ref(verses))

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
